"""Provider-neutral, review-first financial import pipeline."""
from __future__ import annotations
import copy, csv, hashlib, io, math, os, re
from datetime import date, datetime
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from core import _client, db, now_utc, require_admin, round2
from financial_services import entity_for_investment, fingerprint, normalized_key, record_imported_transaction

router = APIRouter(tags=["imports"])
MAX_BYTES, SUPPORTED = 10 * 1024 * 1024, {".csv", ".xlsx", ".xlsm"}

class CommitImportRequest(BaseModel):
    row_numbers: list[int] = Field(min_length=1)
    owner: str = Field(default="Self", min_length=1, max_length=120)
    occurred_on: str | None = None

DATE_FIELDS = ("date", "transaction_date", "trade_date", "occurred_on", "date_time")
PNL_FIELDS = ("realized_pnl", "realised_pnl", "realized_p&l", "realised_p&l", "p&l", "pnl", "profit_loss")


def _number(value):
    try:
        number = float(str(value or "0").replace(",", "").replace("₹", "").strip())
        return round2(number) if math.isfinite(number) else 0.0
    except (ValueError, TypeError): return 0.0


def _date_value(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    text = str(value or "").strip()
    if not text: return None
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try: return datetime.strptime(text[:10], pattern).date().isoformat()
        except ValueError: pass
    return None


def _clean_row(row):
    return {str(key or "").strip().lower().replace(" ", "_"): value for key, value in row.items() if str(key or "").strip()}

def detect_provider(filename, headers):
    signal = f"{filename} {' '.join(headers)}".lower()
    if "zerodha" in signal or "tradingsymbol" in signal or "isin" in signal: return "ZERODHA", "BROKER_EXPORT", .94
    if "groww" in signal: return "GROWW", "BROKER_EXPORT", .90
    if "upstox" in signal: return "UPSTOX", "BROKER_EXPORT", .90
    return "GENERIC", "TABULAR_FINANCIAL_EXPORT", .58

def classify_row(row):
    clean = _clean_row(row); text = " ".join(str(value or "") for value in clean.values()).lower()
    name = clean.get("symbol") or clean.get("tradingsymbol") or clean.get("instrument") or clean.get("scrip") or clean.get("name") or "Imported investment"
    pnl = next((_number(clean.get(key)) for key in ("realized_pnl", "realised_pnl", "realized_p&l", "realised_p&l", "p&l", "pnl", "profit_loss") if clean.get(key) not in (None, "")), None)
    quantity = next((_number(clean.get(key)) for key in ("quantity", "qty", "net_quantity")), 0)
    value = next((_number(clean.get(key)) for key in ("current_value", "market_value", "value", "closing_value", "amount")), 0)
    cost = next((_number(clean.get(key)) for key in ("invested", "cost", "buy_value", "average_price")), 0)
    isin = str(clean.get("isin") or "").strip()
    if pnl is not None or "p&l" in text or "profit" in text or "loss" in text:
        return {"kind":"HISTORICAL_PNL", "confidence":.91, "name":str(name), "isin":isin, "amount":pnl or 0, "reason":"Recognised a realised P&L field; it will not be treated as a current holding."}
    if name and (quantity or value or cost):
        return {"kind":"HOLDING", "confidence":.86, "name":str(name), "isin":isin, "quantity":quantity, "current_value":value or cost, "cost":cost or value, "reason":"Recognised an instrument with holding, quantity, cost or valuation fields."}
    return {"kind":"NEEDS_GUIDANCE", "confidence":.35, "name":str(name), "isin":isin, "reason":"No safe financial interpretation was found. Choose a mapping in Review Inbox."}

def parse_rows(filename, content):
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".csv": return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
    if ext in {".xlsx", ".xlsm"}:
        values = list(load_workbook(io.BytesIO(content), read_only=True, data_only=True).active.iter_rows(values_only=True))
        if not values: return []
        headers = [str(value or "").strip() for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(value not in (None, "") for value in row)]
    raise HTTPException(status_code=415, detail="This file type needs a parser adapter. CSV, XLSX and XLSM are currently supported for safe extraction.")


def _preview_candidate(source, candidate):
    clean = _clean_row(source)
    errors = []
    raw_pnl = next((clean.get(key) for key in PNL_FIELDS if clean.get(key) not in (None, "")), None)
    if candidate["kind"] == "HISTORICAL_PNL" and raw_pnl is not None:
        try:
            amount = float(str(raw_pnl).replace(",", "").replace("₹", "").strip())
            if not math.isfinite(amount): raise ValueError("non-finite amount")
        except (TypeError, ValueError): errors.append("P&L amount is not a valid number.")
    raw_date = next((clean.get(key) for key in DATE_FIELDS if clean.get(key) not in (None, "")), None)
    occurred_on = _date_value(raw_date)
    if raw_date is not None and occurred_on is None:
        errors.append("Transaction date is invalid; use YYYY-MM-DD or DD/MM/YYYY.")
    if occurred_on: candidate["occurred_on"] = occurred_on
    candidate["errors"] = errors
    candidate["can_apply"] = candidate["kind"] != "NEEDS_GUIDANCE" and not errors
    candidate["planned_change"] = {
        "HOLDING": "Create or update this investment holding",
        "HISTORICAL_PNL": "Add a posted historical Investment P&L transaction",
    }.get(candidate["kind"], "No change until the row is mapped")
    return candidate


async def _likely_existing_transaction(candidate):
    """Flag exact title/type/amount matches without silently treating them as duplicates."""
    if candidate.get("kind") != "HISTORICAL_PNL":
        return None
    amount = abs(round2(candidate.get("amount", 0)))
    transaction_type = "INCOME" if candidate.get("amount", 0) >= 0 else "EXPENSE"
    title = re.escape(str(candidate.get("name") or "").strip())
    if not title:
        return None
    query = {
        "type": transaction_type,
        "title": {"$regex": f"^{title}$", "$options": "i"},
        "amount": {"$gte": amount - 0.01, "$lte": amount + 0.01},
        "deleted_at": {"$exists": False},
    }
    if candidate.get("occurred_on"):
        query["date"] = candidate["occurred_on"]
    return await db.transactions.find_one(query)


def _candidate_fingerprint(provider, candidate):
    parts = [
        provider, candidate.get("isin") or candidate["name"], candidate["kind"],
        candidate.get("quantity", candidate.get("amount", 0)),
        candidate.get("current_value", candidate.get("amount", 0)),
    ]
    if candidate["kind"] == "HISTORICAL_PNL":
        parts.append(candidate.get("occurred_on", ""))
    return fingerprint(*parts)


@router.post("/imports/analyze")
async def analyze_import(file: UploadFile = File(...), user: dict = Depends(require_admin)):
    filename, content = file.filename or "upload", await file.read()
    if not content: raise HTTPException(status_code=400, detail="The uploaded file is empty")
    if len(content) > MAX_BYTES: raise HTTPException(status_code=413, detail="Import file is too large (10MB maximum)")
    content_sha256 = hashlib.sha256(content).hexdigest(); previous = await db.import_runs.find_one({"content_sha256": content_sha256})
    if previous: return {"status":"ALREADY_IMPORTED", "run_id":str(previous["_id"]), "filename":previous.get("filename"), "message":"This exact document was already analyzed. Open its audit record instead of importing it again."}
    rows = parse_rows(filename, content)
    if not rows: raise HTTPException(status_code=400, detail="No tabular records were found")
    provider, document_type, confidence = detect_provider(filename, list(_clean_row(rows[0]).keys())); now = now_utc()
    run = {"filename":filename,"content_sha256":content_sha256,"provider":provider,"document_type":document_type,"document_confidence":confidence,"status":"REVIEW","created_at":now,"created_by":user["email"],"row_count":len(rows),"parser_version":"2026.09.20"}
    result = await db.import_runs.insert_one(run); run_id = str(result.inserted_id); candidates=[]
    fingerprints = {}
    preview_errors = 0
    for number, source in enumerate(rows[:3000], start=2):
        candidate = _preview_candidate(source, classify_row(source))
        row_fp = _candidate_fingerprint(provider, candidate)
        committed = await db.import_rows.find_one({"fingerprint":row_fp,"status":"COMMITTED"})
        earlier = fingerprints.get(row_fp)
        existing_transaction = await _likely_existing_transaction(candidate)
        duplicate = bool(committed or earlier is not None or existing_transaction)
        candidate.update({
            "row_number":number,
            "fingerprint":row_fp,
            "duplicate":duplicate,
            "duplicate_of_row":earlier,
            "duplicate_reason":(
                f"Matches row {earlier} in this file." if earlier is not None
                else "Matches a previously imported row." if committed
                else "A transaction with the same title, type and amount already exists." if existing_transaction
                else None
            ),
            "source":source,
        })
        if candidate["errors"]: preview_errors += len(candidate["errors"])
        if not duplicate: fingerprints[row_fp] = number
        await db.import_rows.insert_one({
            "run_id":run_id,"row_number":number,"fingerprint":row_fp,"candidate":candidate,
            "source_row":source,"status":"DUPLICATE" if duplicate else "REVIEW","created_at":now,
        })
        candidates.append(candidate)
    return {
        "status":"REVIEW","run_id":run_id,"filename":filename,"provider":provider,
        "document_type":document_type,"document_confidence":confidence,"total_rows":len(rows),
        "candidates":candidates,
        "preview":{"ready":sum(row["can_apply"] and not row["duplicate"] for row in candidates),
                   "duplicates":sum(row["duplicate"] for row in candidates),
                   "errors":preview_errors,
                   "needs_guidance":sum(row["kind"] == "NEEDS_GUIDANCE" for row in candidates),
                   "not_previewed":max(0, len(rows) - len(candidates))},
        "notice":"Nothing has changed in your finances. Review the proposed changes and row errors before approving.",
    }

async def _commit_row(row, run, request, session=None):
    candidate, kind = row["candidate"], row["candidate"]["kind"]
    if kind == "NEEDS_GUIDANCE": return "NEEDS_GUIDANCE"
    if row.get("status") == "DUPLICATE" or candidate.get("duplicate"): return "DUPLICATE"
    if candidate.get("errors"): raise ValueError(f"Row {row['row_number']} has preview errors and cannot be applied")
    owner = request.owner.strip(); occurrence = request.occurred_on or candidate.get("occurred_on") or now_utc().date().isoformat()
    entity_key = normalized_key("investment", candidate["name"], owner, "Equity")
    previous_entity = await db.financial_entities.find_one({"identity_key":entity_key}, session=session)
    entity = await entity_for_investment(candidate["name"], owner, "Equity", str(run["_id"]), session=session)
    audited_documents = []
    if kind == "HOLDING":
        identity = normalized_key(candidate.get("isin") or candidate["name"], owner, "Equity"); existing = await db.investments.find_one({"identity_key":identity,"deleted_at":{"$exists":False}}, session=session)
        previous_holding = await db.holdings.find_one({"identity_key":identity}, session=session)
        update = {"name":candidate["name"],"symbol":candidate["name"],"isin":candidate.get("isin"),"type":"Stocks","asset_class":"Equity","cost":candidate.get("cost",0),"current_value":candidate.get("current_value",0),"quantity":candidate.get("quantity",0),"owner":owner,"ownership_percent":100,"identity_key":identity,"canonical_entity_id":str(entity["_id"]),"source_import_id":str(run["_id"]),"source_document":run["filename"],"source_row":row["row_number"],"updated_at":now_utc()}
        if existing: await db.investments.update_one({"_id":existing["_id"]},{"$set":update}, session=session); outcome="CONSOLIDATED"
        else: update["created_at"]=now_utc(); await db.investments.insert_one(update, session=session); outcome="CREATED"
        await db.holdings.update_one({"identity_key":identity},{"$set":{"entity_id":str(entity["_id"]),"identity_key":identity,"symbol":candidate["name"],"isin":candidate.get("isin"),"quantity":candidate.get("quantity",0),"cost":candidate.get("cost",0),"market_value":candidate.get("current_value",0),"valuation_date":occurrence,"source_import_id":str(run["_id"]),"source_row":row["row_number"],"updated_at":now_utc()},"$setOnInsert":{"created_at":now_utc()}}, upsert=True, session=session)
        current_investment = await db.investments.find_one({"identity_key":identity,"source_import_id":str(run["_id"]),"source_row":row["row_number"]}, session=session)
        current_holding = await db.holdings.find_one({"identity_key":identity}, session=session)
        audited_documents.extend([
            {"collection":"investments","id":current_investment["_id"],"before":copy.deepcopy(existing),"after":copy.deepcopy(current_investment)},
            {"collection":"holdings","id":current_holding["_id"],"before":copy.deepcopy(previous_holding),"after":copy.deepcopy(current_holding)},
        ])
        await record_imported_transaction(run_id=str(run["_id"]),row_number=row["row_number"],entity_id=str(entity["_id"]),kind="HOLDING_VALUATION",amount=candidate.get("current_value",0),occurred_on=occurrence,source_reference=run["filename"], session=session)
    else:
        amount=candidate.get("amount",0)
        await record_imported_transaction(run_id=str(run["_id"]),row_number=row["row_number"],entity_id=str(entity["_id"]),kind="HISTORICAL_PNL",amount=amount,occurred_on=occurrence,source_reference=run["filename"], session=session)
        await db.transactions.insert_one({"type":"INCOME" if amount>=0 else "EXPENSE","amount":abs(amount),"category":"Investment P&L","source":run["provider"],"title":candidate["name"],"description":"Reviewed historical broker P&L import","date":occurrence,"scope":"PERSONAL","transaction_status":"POSTED","record_source":"IMPORT","source_import_id":str(run["_id"]),"source_document":run["filename"],"source_row":row["row_number"],"created_at":now_utc(),"created_by":run["created_by"]}, session=session); outcome="CREATED"
        imported_transaction = await db.transactions.find_one({"source_import_id":str(run["_id"]),"source_row":row["row_number"]}, session=session)
        if imported_transaction:
            audited_documents.append({"collection":"transactions","id":imported_transaction["_id"],"before":None,"after":copy.deepcopy(imported_transaction)})
    canonical_transaction = await db.financial_transactions.find_one(
        {"source_import_id":str(run["_id"]),"source_row":row["row_number"]}, session=session
    )
    if canonical_transaction:
        audited_documents.append({"collection":"financial_transactions","id":canonical_transaction["_id"],"before":None,"after":copy.deepcopy(canonical_transaction)})
    undo_record = {"version":1,"documents":audited_documents}
    if previous_entity is None:
        current_entity = await db.financial_entities.find_one({"_id":entity["_id"]}, session=session)
        undo_record["created_entity"] = {"id":entity["_id"],"after":copy.deepcopy(current_entity)}
    await db.import_rows.update_one(
        {"_id":row["_id"]},
        {"$set":{"status":"COMMITTED","committed_at":now_utc(),"outcome":outcome,"undo_record":undo_record}},
        session=session,
    )
    return outcome


def _transaction_unavailable(exc):
    message = str(exc).lower()
    return "transaction numbers are only allowed" in message or "replica set" in message or "transactions are not supported" in message


@router.post("/imports/{run_id}/commit")
async def commit_import(run_id: str, request: CommitImportRequest, user: dict = Depends(require_admin)):
    from bson import ObjectId
    try: object_id=ObjectId(run_id)
    except Exception as exc: raise HTTPException(status_code=400, detail="Invalid import run") from exc
    run=await db.import_runs.find_one({"_id":object_id,"status":"REVIEW"})
    if not run: raise HTTPException(status_code=404, detail="Import review was not found or has already been committed")
    selected=set(request.row_numbers); rows=await db.import_rows.find({"run_id":run_id,"row_number":{"$in":list(selected)}}).to_list(3000)
    if len(rows)!=len(selected): raise HTTPException(status_code=400, detail="One or more selected rows do not belong to this import")
    rows.sort(key=lambda row: row["row_number"])
    if any(row.get("candidate", {}).get("errors") for row in rows):
        raise HTTPException(status_code=400, detail="Selected rows contain preview errors. Correct or exclude those rows before applying the import.")
    summary={"created":0,"consolidated":0,"duplicates":0,"needs_guidance":0}; keys={"CREATED":"created","CONSOLIDATED":"consolidated","DUPLICATE":"duplicates","NEEDS_GUIDANCE":"needs_guidance"}
    try:
        async with await _client.start_session() as session:
            async with session.start_transaction():
                for row in rows: summary[keys[await _commit_row(row,run,request,session)]]+=1
    except Exception as exc:
        if not _transaction_unavailable(exc): raise
        for row in rows: summary[keys[await _commit_row(row,run,request)]]+=1
    await db.import_runs.update_one({"_id":object_id},{"$set":{"status":"COMMITTED","committed_at":now_utc(),"committed_by":user["email"],"summary":summary}})
    return {"run_id":run_id,"summary":summary}


async def _undo_document(entry, session=None):
    from bson import ObjectId

    collection = getattr(db, entry["collection"])
    identifier = entry["id"]
    if isinstance(identifier, str):
        identifier = ObjectId(identifier)
    current = await collection.find_one({"_id":identifier}, session=session)
    before, after = entry.get("before"), entry["after"]
    if current is None:
        if before is None: return
        raise ValueError(f"Imported {entry['collection']} record {identifier} is missing; it was not restored.")
    if before is not None and current == before:
        return
    if current != after:
        raise ValueError(f"Imported {entry['collection']} record {identifier} changed after import; left it untouched.")
    if before is None:
        await collection.delete_one({"_id":identifier}, session=session)
    else:
        await collection.replace_one({"_id":identifier}, before, session=session)


async def _undo_row(row, session=None):
    if row.get("status") == "UNDONE":
        return {"row_number":row["row_number"],"status":"UNDONE"}
    undo_record = row.get("undo_record")
    if not undo_record:
        raise ValueError("No undo audit snapshot exists for this row; no changes were made.")
    for entry in reversed(undo_record.get("documents", [])):
        await _undo_document(entry, session=session)
    retained_entity = False
    entity = undo_record.get("created_entity")
    if entity:
        from bson import ObjectId
        identifier = entity["id"]
        if isinstance(identifier, str): identifier = ObjectId(identifier)
        current = await db.financial_entities.find_one({"_id":identifier}, session=session)
        if current is not None:
            if current != entity["after"]:
                raise ValueError(f"Imported financial entity {identifier} changed after import; left it untouched.")
            referenced = (
                await db.financial_transactions.find_one({"entity_id":str(identifier)}, session=session)
                or await db.holdings.find_one({"entity_id":str(identifier)}, session=session)
            )
            if referenced:
                retained_entity = True
            else:
                await db.financial_entities.delete_one({"_id":identifier}, session=session)
    result = {"row_number":row["row_number"],"status":"UNDONE"}
    if retained_entity: result["note"] = "A shared financial entity was retained because another record still references it."
    await db.import_rows.update_one(
        {"_id":row["_id"]},
        {"$set":{"status":"UNDONE","undone_at":now_utc(),"undo_result":result},"$unset":{"undo_error":""}},
        session=session,
    )
    return result


@router.post("/imports/{run_id}/undo")
async def undo_import(run_id: str, user: dict = Depends(require_admin)):
    from bson import ObjectId
    try: object_id=ObjectId(run_id)
    except Exception as exc: raise HTTPException(status_code=400, detail="Invalid import run") from exc
    run = await db.import_runs.find_one({"_id":object_id})
    if not run: raise HTTPException(status_code=404, detail="Import batch was not found")
    if run.get("status") not in {"COMMITTED","UNDO_PARTIAL"}:
        raise HTTPException(status_code=409, detail="Only an applied import batch can be undone.")
    rows = await db.import_rows.find({"run_id":run_id,"status":{"$in":["COMMITTED","UNDO_FAILED"]}}).sort("row_number",-1).to_list(3000)
    completed, failures = [], []
    fallback = False
    for row in rows:
        try:
            if fallback:
                completed.append(await _undo_row(row))
                continue
            async with await _client.start_session() as session:
                async with session.start_transaction():
                    completed.append(await _undo_row(row, session=session))
        except Exception as exc:
            if _transaction_unavailable(exc):
                fallback = True
                try:
                    completed.append(await _undo_row(row))
                except Exception as fallback_exc:
                    failures.append({"row_number":row["row_number"],"error":str(fallback_exc)})
                    await db.import_rows.update_one({"_id":row["_id"]},{"$set":{"status":"UNDO_FAILED","undo_error":str(fallback_exc)}})
            else:
                failures.append({"row_number":row["row_number"],"error":str(exc)})
                await db.import_rows.update_one({"_id":row["_id"]},{"$set":{"status":"UNDO_FAILED","undo_error":str(exc)}})
    status = "UNDO_PARTIAL" if failures else "UNDONE"
    await db.import_runs.update_one({"_id":object_id},{"$set":{
        "status":status,"undo_status":status,"undone_at":now_utc(),"undone_by":user["email"],
        "undo_failures":failures,
    }})
    return {"run_id":run_id,"status":status,"undone_rows":completed,"failures":failures,
            "message":"Some rows could not be safely undone; they were left unchanged and are listed for review." if failures else "All applied rows were undone."}


@router.get("/imports")
async def list_imports(user: dict = Depends(require_admin)):
    rows=await db.import_runs.find({}).sort("created_at",-1).to_list(200)
    return [{"id":str(row["_id"]),"filename":row.get("filename"),"provider":row.get("provider"),"document_type":row.get("document_type"),"status":row.get("status"),"row_count":row.get("row_count"),"created_at":row.get("created_at"),"summary":row.get("summary",{})} for row in rows]
